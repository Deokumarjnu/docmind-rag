"""Document loaders for various file formats.

This module provides unified loading for multiple document types:
- PDF, DOCX, HTML, TXT (original)
- CSV, XLSX/XLS (spreadsheets)
- JSON, JSONL (structured data)
- Markdown (documentation)
"""

import json
import logging
from pathlib import Path
from typing import Optional
import pandas as pd

from langchain_core.documents import Document

logger = logging.getLogger(__name__)


class UnifiedDocumentLoader:
    """
    Unified document loader that handles multiple file formats.
    
    Routes to the appropriate loader based on file extension.
    """
    
    # File extension to loader mapping
    LOADER_MAP = {
        # Text-based
        "txt": "text",
        "html": "html",
        "htm": "html",
        
        # Documents
        "pdf": "pdf",
        "docx": "docx",
        "doc": "docx",
        
        # Spreadsheets
        "csv": "csv",
        "xlsx": "excel",
        "xls": "excel",
        
        # Structured data
        "json": "json",
        "jsonl": "jsonl",
        
        # Markdown
        "md": "markdown",
        "markdown": "markdown",
    }
    
    def __init__(self):
        """Initialize the unified loader."""
        self._check_dependencies()
    
    def _check_dependencies(self):
        """Check for optional dependencies."""
        self._has_pandas = False
        self._has_openpyxl = False
        
        try:
            import pandas
            self._has_pandas = True
        except ImportError:
            logger.warning("pandas not installed - CSV/Excel loading may be limited")
        
        try:
            import openpyxl
            self._has_openpyxl = True
        except ImportError:
            logger.warning("openpyxl not installed - Excel loading may be limited")
    
    def get_file_type(self, file_path: str | Path) -> str:
        """Get the file type from extension."""
        ext = Path(file_path).suffix.lower().lstrip(".")
        return self.LOADER_MAP.get(ext, "unknown")
    
    def load(self, file_path: str | Path) -> list[Document]:
        """
        Load a document from file path.
        
        Args:
            file_path: Path to the document file
            
        Returns:
            List of Document objects
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_type = self.get_file_type(file_path)
        
        loader_method = {
            "text": self._load_text,
            "html": self._load_html,
            "pdf": self._load_pdf,
            "docx": self._load_docx,
            "csv": self._load_csv,
            "excel": self._load_excel,
            "json": self._load_json,
            "jsonl": self._load_jsonl,
            "markdown": self._load_markdown,
        }.get(file_type)
        
        if loader_method is None:
            raise ValueError(f"Unsupported file type: {file_path.suffix}")
        
        logger.info(f"Loading {file_type} file: {file_path.name}")
        return loader_method(file_path)
    
    def _load_text(self, file_path: Path) -> list[Document]:
        """Load plain text file."""
        try:
            from langchain_community.document_loaders import TextLoader
            loader = TextLoader(str(file_path), encoding="utf-8")
            docs = loader.load()
        except Exception:
            # Fallback to direct reading
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            docs = [Document(
                page_content=content,
                metadata={"source": str(file_path), "file_type": "txt"}
            )]
        
        for doc in docs:
            doc.metadata["file_type"] = "txt"
            doc.metadata["content_type"] = "text"
        
        return docs
    
    def _load_html(self, file_path: Path) -> list[Document]:
        """Load HTML file."""
        try:
            from langchain_community.document_loaders import UnstructuredHTMLLoader
            loader = UnstructuredHTMLLoader(str(file_path))
            docs = loader.load()
        except ImportError:
            # Fallback: strip HTML tags with basic regex
            import re
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            # Remove script and style elements
            content = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", content, flags=re.DOTALL | re.IGNORECASE)
            # Remove HTML tags
            content = re.sub(r"<[^>]+>", " ", content)
            # Clean up whitespace
            content = re.sub(r"\s+", " ", content).strip()
            docs = [Document(
                page_content=content,
                metadata={"source": str(file_path), "file_type": "html"}
            )]
        
        for doc in docs:
            doc.metadata["file_type"] = "html"
            doc.metadata["content_type"] = "text"
        
        return docs
    
    def _load_pdf(self, file_path: Path) -> list[Document]:
        """Load PDF file using PyMuPDF."""
        try:
            from langchain_community.document_loaders import PyMuPDFLoader
            loader = PyMuPDFLoader(str(file_path))
            docs = loader.load()
        except Exception as e:
            logger.error(f"Failed to load PDF: {e}")
            raise
        
        for doc in docs:
            doc.metadata["file_type"] = "pdf"
        
        return docs
    
    def _load_docx(self, file_path: Path) -> list[Document]:
        """Load DOCX file."""
        try:
            from langchain_community.document_loaders import Docx2txtLoader
            loader = Docx2txtLoader(str(file_path))
            docs = loader.load()
        except ImportError:
            try:
                from langchain_community.document_loaders import UnstructuredWordDocumentLoader
                loader = UnstructuredWordDocumentLoader(str(file_path))
                docs = loader.load()
            except Exception as e:
                logger.error(f"Failed to load DOCX (install docx2txt or unstructured): {e}")
                raise
        
        for doc in docs:
            doc.metadata["file_type"] = "docx"
            doc.metadata["content_type"] = "text"
        
        return docs
    
    def _load_csv(self, file_path: Path) -> list[Document]:
        """
        Load CSV file.
        
        Converts CSV to a structured text format that's better for RAG.
        """
        try:
            from langchain_community.document_loaders import CSVLoader
            loader = CSVLoader(str(file_path), encoding="utf-8")
            docs = loader.load()
        except ImportError:
            # Fallback using pandas
            if self._has_pandas:
                import pandas as pd
                df = pd.read_csv(file_path)
                docs = self._dataframe_to_documents(df, file_path, "csv")
            else:
                # Basic CSV reading
                import csv
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    reader = csv.DictReader(f)
                    rows = list(reader)
                
                docs = []
                for i, row in enumerate(rows):
                    content = "\n".join([f"{k}: {v}" for k, v in row.items() if v])
                    docs.append(Document(
                        page_content=content,
                        metadata={
                            "source": str(file_path),
                            "row": i,
                            "file_type": "csv",
                            "content_type": "table",
                        }
                    ))
        
        for doc in docs:
            doc.metadata["file_type"] = "csv"
            doc.metadata["content_type"] = "table"
        
        logger.info(f"Loaded CSV with {len(docs)} rows")
        return docs
    
    def _load_excel(self, file_path: Path) -> list[Document]:
        """
        Load Excel file (XLSX/XLS).
        
        Processes each sheet as a separate document section.
        """
        if not self._has_pandas:
            raise ImportError("pandas is required for Excel loading. Install with: pip install pandas openpyxl")
        
        import pandas as pd
        
        docs = []
        
        try:
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                
                if df.empty:
                    continue
                
                # Convert sheet to documents
                sheet_docs = self._dataframe_to_documents(
                    df, 
                    file_path, 
                    "excel",
                    sheet_name=sheet_name
                )
                docs.extend(sheet_docs)
            
            logger.info(f"Loaded Excel with {len(excel_file.sheet_names)} sheets, {len(docs)} documents")
            
        except Exception as e:
            logger.error(f"Failed to load Excel: {e}")
            raise
        
        return docs
    
    def _dataframe_to_documents(
        self, 
        df, 
        file_path: Path, 
        file_type: str,
        sheet_name: Optional[str] = None,
    ) -> list[Document]:
        """
        Convert a pandas DataFrame to Document objects.
        
        Creates documents in two formats:
        1. Summary document with column info and stats
        2. Row-by-row documents for detailed retrieval
        """
        
        docs = []
        
        # Create summary document
        summary_parts = [
            f"# {'Sheet: ' + sheet_name if sheet_name else 'Data Summary'}",
            f"Rows: {len(df)}, Columns: {len(df.columns)}",
            "",
            "## Columns:",
        ]
        
        for col in df.columns:
            dtype = str(df[col].dtype)
            non_null = df[col].notna().sum()
            summary_parts.append(f"- {col} ({dtype}): {non_null} values")
        
        # Add sample data
        summary_parts.extend([
            "",
            "## Sample Data (first 5 rows):",
            df.head().to_string(),
        ])
        
        docs.append(Document(
            page_content="\n".join(summary_parts),
            metadata={
                "source": str(file_path),
                "file_type": file_type,
                "content_type": "table",
                "sheet_name": sheet_name or "default",
                "document_section": "summary",
                "row_count": len(df),
                "column_count": len(df.columns),
            }
        ))
        
        # Create row documents (for detailed retrieval)
        # Create ONE document per row for better semantic search
        # This allows queries like "Leo Garcia salary" to find the exact row
        for idx, row in df.iterrows():
            # Convert row to readable format with clear field labels
            row_parts = []
            for col, val in row.items():
                if pd.notna(val):
                    row_parts.append(f"{col}: {val}")
            
            row_text = "\n".join(row_parts)
            
            docs.append(Document(
                page_content=row_text,
                metadata={
                    "source": str(file_path),
                    "file_type": file_type,
                    "content_type": "table",
                    "sheet_name": sheet_name or "default",
                    "document_section": "data",
                    "row_index": idx,
                }
            ))
        
        return docs
    
    def _load_json(self, file_path: Path) -> list[Document]:
        """
        Load JSON file.
        
        Handles both single objects and arrays of objects.
        """
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        
        docs = self._json_to_documents(data, file_path)
        
        for doc in docs:
            doc.metadata["file_type"] = "json"
        
        logger.info(f"Loaded JSON with {len(docs)} documents")
        return docs
    
    def _load_jsonl(self, file_path: Path) -> list[Document]:
        """Load JSON Lines file (one JSON object per line)."""
        docs = []
        
        with open(file_path, "r", encoding="utf-8") as f:
            for i, line in enumerate(f):
                line = line.strip()
                if not line:
                    continue
                
                try:
                    obj = json.loads(line)
                    content = self._json_object_to_text(obj)
                    docs.append(Document(
                        page_content=content,
                        metadata={
                            "source": str(file_path),
                            "file_type": "jsonl",
                            "content_type": "structured",
                            "line_number": i + 1,
                        }
                    ))
                except json.JSONDecodeError as e:
                    logger.warning(f"Skipping invalid JSON at line {i + 1}: {e}")
        
        logger.info(f"Loaded JSONL with {len(docs)} lines")
        return docs
    
    def _json_to_documents(self, data, file_path: Path) -> list[Document]:
        """Convert JSON data to Document objects."""
        docs = []
        
        if isinstance(data, list):
            # Array of objects
            for i, item in enumerate(data):
                content = self._json_object_to_text(item)
                docs.append(Document(
                    page_content=content,
                    metadata={
                        "source": str(file_path),
                        "content_type": "structured",
                        "array_index": i,
                    }
                ))
        elif isinstance(data, dict):
            # Single object or nested structure
            content = self._json_object_to_text(data)
            docs.append(Document(
                page_content=content,
                metadata={
                    "source": str(file_path),
                    "content_type": "structured",
                }
            ))
        else:
            # Primitive value
            docs.append(Document(
                page_content=str(data),
                metadata={
                    "source": str(file_path),
                    "content_type": "text",
                }
            ))
        
        return docs
    
    def _json_object_to_text(self, obj, prefix: str = "") -> str:
        """
        Convert a JSON object to readable text format.
        
        Flattens nested structures for better retrieval.
        """
        if not isinstance(obj, dict):
            return str(obj)
        
        lines = []
        
        for key, value in obj.items():
            full_key = f"{prefix}.{key}" if prefix else key
            
            if isinstance(value, dict):
                # Recurse into nested object
                nested_text = self._json_object_to_text(value, full_key)
                lines.append(nested_text)
            elif isinstance(value, list):
                if all(isinstance(item, dict) for item in value):
                    # List of objects
                    for i, item in enumerate(value):
                        item_text = self._json_object_to_text(item, f"{full_key}[{i}]")
                        lines.append(item_text)
                else:
                    # List of primitives
                    lines.append(f"{full_key}: {', '.join(str(v) for v in value)}")
            else:
                lines.append(f"{full_key}: {value}")
        
        return "\n".join(lines)
    
    def _load_markdown(self, file_path: Path) -> list[Document]:
        """
        Load Markdown file.
        
        Preserves structure and splits by headers for better retrieval.
        """
        try:
            from langchain_community.document_loaders import UnstructuredMarkdownLoader
            loader = UnstructuredMarkdownLoader(str(file_path), mode="elements")
            docs = loader.load()
        except ImportError:
            # Fallback: Load as text with basic markdown handling
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()
            
            # Split by headers for better chunking
            docs = self._split_markdown_by_headers(content, file_path)
        
        for doc in docs:
            doc.metadata["file_type"] = "markdown"
            doc.metadata["content_type"] = "text"
        
        logger.info(f"Loaded Markdown with {len(docs)} sections")
        return docs
    
    def _split_markdown_by_headers(self, content: str, file_path: Path) -> list[Document]:
        """Split markdown content by headers."""
        import re
        
        # Split by headers (##, ###, etc.)
        header_pattern = r'^(#{1,6})\s+(.+)$'
        
        lines = content.split("\n")
        sections = []
        current_section = {"header": "Introduction", "level": 0, "content": []}
        
        for line in lines:
            match = re.match(header_pattern, line)
            if match:
                # Save previous section
                if current_section["content"]:
                    sections.append(current_section)
                
                # Start new section
                current_section = {
                    "header": match.group(2).strip(),
                    "level": len(match.group(1)),
                    "content": [line],
                }
            else:
                current_section["content"].append(line)
        
        # Add last section
        if current_section["content"]:
            sections.append(current_section)
        
        # Convert to documents
        docs = []
        for section in sections:
            content = "\n".join(section["content"]).strip()
            if content:
                docs.append(Document(
                    page_content=content,
                    metadata={
                        "source": str(file_path),
                        "header": section["header"],
                        "header_level": section["level"],
                    }
                ))
        
        return docs if docs else [Document(
            page_content=content,
            metadata={"source": str(file_path)}
        )]


# Convenience function for single-file loading
def load_document(file_path: str | Path) -> list[Document]:
    """
    Load a document from file path.
    
    Args:
        file_path: Path to the document file
        
    Returns:
        List of Document objects
    """
    loader = UnifiedDocumentLoader()
    return loader.load(file_path)


def get_supported_extensions() -> list[str]:
    """Get list of supported file extensions."""
    return list(UnifiedDocumentLoader.LOADER_MAP.keys())

